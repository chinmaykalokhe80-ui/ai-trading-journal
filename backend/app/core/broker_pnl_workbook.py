"""Read Zerodha-style P&L workbooks without treating summary rows as trades."""

from datetime import date, datetime
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile

import openpyxl
import pandas as pd


TRADE_HEADERS = ('symbol', 'quantity', 'buy value', 'sell value', 'realized p&l')
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_SHEET_CELLS = 250_000


def label(value):
    return ' '.join(str(value or '').strip().lower().split())


def amount(value, field):
    if value is None or value == '':
        return 0.0
    if isinstance(value, str):
        value = value.strip().replace(',', '').replace('₹', '')
        if value.startswith('(') and value.endswith(')'):
            value = '-' + value[1:-1]
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'Invalid {field} value in workbook: {value!r}') from exc
    if not (-float('inf') < result < float('inf')):
        raise ValueError(f'Invalid {field} value in workbook.')
    return result


def rounded(value):
    return round(value, 2)


def cell_rows(sheet):
    if sheet.max_row * sheet.max_column > MAX_SHEET_CELLS:
        raise ValueError('Workbook sheet is too large to analyze.')
    return list(sheet.values)


def header_at(rows, required):
    for index, row in enumerate(rows):
        columns = {label(value): position for position, value in enumerate(row) if label(value)}
        if all(name in columns for name in required):
            return index, columns
    return None, None


def value_at(row, columns, key):
    index = columns.get(key)
    return row[index] if index is not None and index < len(row) else None


def pair_values(rows, start, end):
    pairs = {}
    for row in rows[start:end]:
        if len(row) > 2 and isinstance(row[1], str) and row[1].strip() and row[2] is not None:
            pairs[label(row[1])] = row[2]
    return pairs


def period_from(rows):
    for row in rows:
        for value in row:
            if not isinstance(value, str):
                continue
            match = re.search(r'\bfrom\s+(\d{4}-\d{2}-\d{2})\s+to\s+(\d{4}-\d{2}-\d{2})\b', value, re.I)
            if match:
                return {'from': match.group(1), 'to': match.group(2)}
    return None


def contract_identity(symbol):
    upper = symbol.upper().replace(' ', '')
    option_type = 'CE' if upper.endswith('CE') else 'PE' if upper.endswith('PE') else 'FUT' if upper.endswith('FUT') else 'Other'
    underlying = next((name for name in ('MIDCPNIFTY', 'BANKNIFTY', 'FINNIFTY', 'SENSEX', 'NIFTY')
                       if upper.startswith(name)), None)
    if underlying is None:
        match = re.match(r'^[A-Z]+', upper)
        underlying = match.group(0) if match else 'Unknown'
    return underlying, option_type


def contract_insights(contracts):
    totals = {
        'contract_rows': len(contracts),
        'quantity': rounded(sum(item['Quantity'] for item in contracts)),
        'buy_value': rounded(sum(item['Buy Value'] for item in contracts)),
        'sell_value': rounded(sum(item['Sell Value'] for item in contracts)),
        'realized_pnl': rounded(sum(item['Realized P&L'] for item in contracts)),
        'unrealized_pnl': rounded(sum(item['Unrealized P&L'] for item in contracts)),
        'open_contract_rows': sum(item['Open Quantity'] != 0 for item in contracts),
        'open_quantity': rounded(sum(item['Open Quantity'] for item in contracts)),
        'open_value': rounded(sum(item['Open Value'] for item in contracts)),
    }
    breakdowns = {}
    for field in ('underlying', 'option_type'):
        groups = {}
        for item in contracts:
            name = item[field]
            group = groups.setdefault(name, {'name': name, 'contract_rows': 0, 'realized_pnl': 0.0,
                                             'unrealized_pnl': 0.0, 'buy_value': 0.0, 'sell_value': 0.0})
            group['contract_rows'] += 1
            for key, source in (('realized_pnl', 'Realized P&L'), ('unrealized_pnl', 'Unrealized P&L'),
                                ('buy_value', 'Buy Value'), ('sell_value', 'Sell Value')):
                group[key] += item[source]
        breakdowns[field] = sorted(
            [{key: rounded(value) if isinstance(value, float) else value for key, value in group.items()}
             for group in groups.values()], key=lambda group: (-abs(group['realized_pnl']), group['name']))
    return totals, breakdowns


def parse_broker_pnl_workbook(content):
    """Return (contract dataframe, statement metadata), or None for another layout."""
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError('Workbook expands beyond the analysis size limit.')
    except BadZipFile as exc:
        raise ValueError('Invalid Excel workbook.') from exc

    # Some broker exports declare dimension A1 despite containing many populated cells.
    # Normal mode reads the actual worksheet XML; read_only mode loses those rows.
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=False, data_only=True)
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ValueError('Invalid Excel workbook.') from exc

    try:
        for sheet in workbook.worksheets:
            rows = cell_rows(sheet)
            header_index, columns = header_at(rows, TRADE_HEADERS)
            if header_index is None:
                continue

            summary = pair_values(rows, 0, header_index)
            contracts = []
            for row in rows[header_index + 1:]:
                symbol = value_at(row, columns, 'symbol')
                if symbol is None or not str(symbol).strip():
                    continue
                if label(symbol) in ('total', 'grand total', 'subtotal'):
                    continue
                contract = {'Symbol': str(symbol).strip(), 'ISIN': str(value_at(row, columns, 'isin') or '').strip()}
                for key, field in (
                    ('quantity', 'Quantity'), ('buy value', 'Buy Value'),
                    ('sell value', 'Sell Value'), ('realized p&l', 'Realized P&L'),
                    ('realized p&l pct.', 'Realized P&L Pct.'),
                    ('previous closing price', 'Previous Closing Price'),
                    ('open quantity', 'Open Quantity'), ('open value', 'Open Value'),
                    ('unrealized p&l', 'Unrealized P&L'),
                    ('unrealized p&l pct.', 'Unrealized P&L Pct.'),
                ):
                    contract[field] = amount(value_at(row, columns, key), field)
                contract['Open Quantity Type'] = str(value_at(row, columns, 'open quantity type') or '').strip()
                contract['underlying'], contract['option_type'] = contract_identity(contract['Symbol'])
                contracts.append(contract)
            if not contracts:
                raise ValueError('The workbook has no contract rows.')

            charges = []
            charge_header, charge_columns = header_at(rows[:header_index], ('account head', 'amount'))
            if charge_header is not None:
                for row in rows[charge_header + 1:header_index]:
                    name = value_at(row, charge_columns, 'account head')
                    if not name or not str(name).strip():
                        continue
                    value = value_at(row, charge_columns, 'amount')
                    if value is not None:
                        charges.append({'name': str(name).strip(), 'amount': amount(value, 'charge')})

            adjustments = []
            for other in workbook.worksheets:
                if other == sheet:
                    continue
                other_rows = cell_rows(other)
                other_header, other_columns = header_at(other_rows, ('particulars', 'posting date', 'debit', 'credit'))
                if other_header is None:
                    continue
                for row in other_rows[other_header + 1:]:
                    particulars = value_at(row, other_columns, 'particulars')
                    if not particulars or not str(particulars).strip():
                        continue
                    posting = value_at(row, other_columns, 'posting date')
                    if isinstance(posting, (date, datetime)):
                        posting = posting.date().isoformat() if isinstance(posting, datetime) else posting.isoformat()
                    adjustments.append({
                        'particulars': str(particulars).strip(), 'posting_date': str(posting or ''),
                        'debit': rounded(amount(value_at(row, other_columns, 'debit'), 'debit')),
                        'credit': rounded(amount(value_at(row, other_columns, 'credit'), 'credit')),
                    })

            reported = {
                'realized_pnl': rounded(amount(summary.get('realized p&l'), 'summary realized P&L')) if 'realized p&l' in summary else None,
                'unrealized_pnl': rounded(amount(summary.get('unrealized p&l'), 'summary unrealized P&L')) if 'unrealized p&l' in summary else None,
                'charges': amount(summary.get('charges'), 'summary charges') if 'charges' in summary else None,
                'other_credit_debit': amount(summary.get('other credit & debit'), 'summary other credit/debit') if 'other credit & debit' in summary else None,
            }
            totals, breakdowns = contract_insights(contracts)
            realized_sum = totals['realized_pnl']
            unrealized_sum = totals['unrealized_pnl']
            charge_sum = round(sum(item['amount'] for item in charges), 6)
            adjustment_sum = round(sum(item['credit'] - item['debit'] for item in adjustments), 6)
            comparisons = {
                'realized_pnl': realized_sum, 'unrealized_pnl': unrealized_sum,
                'charges': charge_sum if charges else None,
                'other_credit_debit': adjustment_sum if adjustments else None,
            }
            reconciliation = {
                key: {'reported': value, 'detail_total': comparisons[key],
                      'difference': round(value - comparisons[key], 6) if value is not None and comparisons[key] is not None else None}
                for key, value in reported.items()
            }
            statement = {
                'sheet': sheet.title, 'period': period_from(rows), 'reported': reported,
                'charges': charges, 'adjustments': adjustments,
                'reconciliation': reconciliation, 'totals': totals,
                'breakdowns': breakdowns, 'contracts': contracts,
            }
            return pd.DataFrame(contracts), statement
        return None
    finally:
        workbook.close()
