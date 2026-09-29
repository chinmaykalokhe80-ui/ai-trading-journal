from io import BytesIO
import re
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


def broker_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'F&O'
    sheet['B11'] = 'P&L Statement for F&O from 2026-09-01 to 2026-09-29'
    sheet['B15'], sheet['C15'] = 'Charges', 12.345
    sheet['B17'], sheet['C17'] = 'Realized P&L', -50
    sheet['B18'], sheet['C18'] = 'Unrealized P&L', 7
    sheet['B23'], sheet['C23'] = 'Account Head', 'Amount'
    sheet['B24'], sheet['C24'] = 'Brokerage - Z', 12.345
    sheet['B38'] = 'Symbol'
    for column, value in enumerate(('Quantity', 'Buy Value', 'Sell Value', 'Realized P&L',
                                     'Realized P&L Pct.', 'Open Quantity', 'Open Quantity Type',
                                     'Open Value', 'Unrealized P&L'), 4):
        sheet.cell(38, column, value)
    for row, values in enumerate((('NIFTY2690123950CE', 65, 100, 150, 50, 50, 0, '', 0, 0),
                                  ('SENSEX2690374700PE', 20, 200, 100, -100, -50, 20, 'Buy', 80, 7)), 39):
        for column, value in zip((2, 4, 5, 6, 7, 8, 9, 10, 11, 12), values):
            sheet.cell(row, column, value)
    sheet['B41'], sheet['G41'] = 'Total', -50
    other = workbook.create_sheet('Other Debits and Credits')
    for column, value in enumerate(('Particulars', 'Posting Date', 'Debit', 'Credit'), 2):
        other.cell(15, column, value)
    for column, value in enumerate(('Adjustment', '2026-09-28', 10, 0), 2):
        other.cell(16, column, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def wrong_dimension(content):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, 'w') as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                data, count = re.subn(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data, count=1)
                assert count == 1
            target.writestr(member, data)
    return output.getvalue()


def test_broker_workbook_extracts_all_sections_without_counting_totals():
    with TestClient(app) as client:
        response = client.post('/api/ai-coach/analyze-csv', files={
            'file': ('statement.xlsx', wrong_dimension(broker_workbook()),
                     'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    report = response.json()['report']
    statement = report['statement']
    assert report['metrics']['count'] == 2
    assert report['metrics']['pnl'] == -50
    assert report['metrics']['max_drawdown'] is None
    assert statement['period'] == {'from': '2026-09-01', 'to': '2026-09-29'}
    assert statement['totals']['open_contract_rows'] == 1
    assert statement['totals']['unrealized_pnl'] == 7
    assert statement['reported']['charges'] == 12.345
    assert statement['charges'] == [{'name': 'Brokerage - Z', 'amount': 12.345}]
    assert statement['reconciliation']['realized_pnl']['difference'] == 0
    assert statement['reconciliation']['charges']['difference'] == 0
    assert statement['adjustments'][0]['debit'] == 10
    assert [group['name'] for group in statement['breakdowns']['underlying']] == ['SENSEX', 'NIFTY']
    assert len(statement['contracts']) == 2


def test_empty_or_invalid_excel_report_is_rejected():
    with TestClient(app) as client:
        response = client.post('/api/ai-coach/analyze-csv', files={'file': ('bad.xlsx', b'not excel')})
    assert response.status_code == 422


def test_simple_excel_report_remains_supported():
    workbook = Workbook()
    workbook.active.append(('Symbol', 'PnL'))
    workbook.active.append(('ABC', 125))
    output = BytesIO()
    workbook.save(output)
    with TestClient(app) as client:
        response = client.post('/api/ai-coach/analyze-csv', files={'file': ('simple.xlsx', output.getvalue())})
    assert response.status_code == 200
    assert response.json()['report']['metrics']['pnl'] == 125
    assert 'statement' not in response.json()['report']
